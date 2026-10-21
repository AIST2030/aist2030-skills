#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
analyze.py —— 老旧课程 PPT 复活 · 第一步「解析」

把一份 PPTX 解析成结构化 JSON（page-content.json），供模型阅读理解每页教学内容与知识结构。
本脚本只做确定性提取，不做任何文本创作、不改动原文件。

用法：
    python analyze.py 老课件.pptx -w work/ -o work/page-content.json

退出码：
    0 成功 | 2 文件不存在/非 pptx | 3 损坏无法解析 | 4 加密或旧格式
    5 空文档(0 页) | 6 全文无文本（不阻塞，仍产出 JSON）
"""

import argparse
import hashlib
import json
import os
import sys

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    sys.stderr.write("[analyze] 缺少 python-pptx，请先 pip install python-pptx\n")
    sys.exit(1)

EXIT_OK = 0
EXIT_BAD_INPUT = 2
EXIT_PARSE_FAILED = 3
EXIT_ENCRYPTED = 4
EXIT_EMPTY = 5
EXIT_NO_TEXT = 6

# OLE 复合文档魔数：加密的 pptx / 老式 .ppt 都会是这个头
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def fail(code, message, hint=""):
    """按统一结构输出错误并以约定退出码退出"""
    sys.stderr.write(json.dumps(
        {"ok": False, "code": code, "message": message, "hint": hint},
        ensure_ascii=False, indent=2) + "\n")
    sys.exit(code)


def sha1_of(path):
    h = hashlib.sha1()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def iter_shapes(shapes):
    """递归展开组合形状（GROUP），保证文本不漏"""
    for sh in shapes:
        try:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                for sub in iter_shapes(sh.shapes):
                    yield sub
                continue
        except Exception:
            pass
        yield sh


def emu_to_pt(v):
    return round(v / 12700.0, 1) if v else None


TOC_HINTS = ("目录", "CONTENTS", "content", "大纲", "agenda", "AGENDA")
SECTION_HINTS = ("壹", "贰", "叁", "肆", "伍", "第一章", "第二章", "第三章", "第一部分", "第二部分")


def guess_role(page_no, total, texts, title, tables, image_count):
    """
    页面角色识别：不同角色在后面写讲稿时处理方式不同
      cover   封面：开场钩子，建立期待
      toc     目录：预告结构，给听众地图
      section 章节扉页：转场衔接，承上启下
      ending  结束页：回顾与行动号召
      image   纯图页：无文本，需结合画面说明，不得臆造
      content 内容页：正常讲解
    """
    if not texts and not tables:
        return "image"
    if page_no == 1:
        return "cover"
    if page_no == total and total > 2:
        return "ending"

    all_text = " ".join(texts) + " " + (title or "")
    low = all_text.lower()
    for h in TOC_HINTS:
        if h.lower() in low:
            return "toc"
    # 章节扉页：文本很短且含章节序号标记
    if len(all_text) <= 60:
        for h in SECTION_HINTS:
            if h in all_text:
                return "section"
        if image_count >= 1 and len(all_text) <= 24:
            return "section"
    return "content"


def analyze(path):
    if not os.path.exists(path):
        fail(EXIT_BAD_INPUT, "文件不存在：%s" % path, "请确认路径是否正确")
    if os.path.splitext(path)[1].lower() != ".pptx":
        fail(EXIT_BAD_INPUT, "仅支持 .pptx：%s" % path,
             "若为 .ppt / .pptm，请先用 WPS 或 PowerPoint 另存为 .pptx")

    with open(path, "rb") as f:
        head = f.read(8)
    if head == OLE_MAGIC:
        fail(EXIT_ENCRYPTED, "文件是加密包或旧式二进制格式，无法解析",
             "请去除密码，或用 WPS 打开后另存为无保护的 .pptx")

    try:
        prs = Presentation(path)
    except Exception as e:
        fail(EXIT_PARSE_FAILED, "文件损坏或无法解析：%s" % e,
             "建议用 WPS 打开后另存为一份新的 .pptx（另存常能修复轻微损坏）")

    slides = list(prs.slides)
    if not slides:
        fail(EXIT_EMPTY, "该文档不含任何幻灯片", "请确认是否拿错了文件")

    font_set = {}
    font_sizes = []
    pages = []
    total_chars = 0
    has_notes = 0
    no_text_pages = 0

    for idx, slide in enumerate(slides, start=1):
        texts = []
        tables = []
        image_count = 0
        levels = []

        title = ""
        try:
            if slide.shapes.title is not None:
                title = slide.shapes.title.text.strip()
        except Exception:
            title = ""

        for sh in iter_shapes(slide.shapes):
            # 图片
            try:
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image_count += 1
            except Exception:
                pass

            # 表格
            try:
                if getattr(sh, "has_table", False):
                    rows = []
                    for r in sh.table.rows:
                        rows.append([c.text.strip() for c in r.cells])
                    if rows:
                        tables.append(rows)
                    continue
            except Exception:
                pass

            # 文本
            try:
                if not getattr(sh, "has_text_frame", False):
                    continue
                tf = sh.text_frame
            except Exception:
                continue

            for para in tf.paragraphs:
                ptext = "".join(run.text for run in para.runs).strip()
                if not ptext:
                    continue
                texts.append(ptext)
                try:
                    levels.append(para.level)
                except Exception:
                    levels.append(0)

                for run in para.runs:
                    fn = run.font.name
                    if fn:
                        font_set[fn] = font_set.get(fn, 0) + 1
                    try:
                        if run.font.size is not None:
                            font_sizes.append(round(run.font.size.pt, 1))
                    except Exception:
                        pass

        # 备注
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ""
        if notes:
            has_notes += 1

        body_chars = sum(len(t) for t in texts)
        page_chars = body_chars + len(title)
        total_chars += page_chars

        # 信息密度
        elements = len(texts) + len(tables) + image_count
        if not texts and not tables:
            no_text_pages += 1
        if page_chars >= 200 or len(tables) > 0 or elements >= 8:
            density = "high"
        elif page_chars < 60 and elements <= 3:
            density = "low"
        else:
            density = "medium"

        # 该页问题
        issues = []
        if not notes:
            issues.append("备注栏为空")
        if not texts and not tables:
            issues.append("无文本内容（疑似纯图片页）")
        if len(tables) > 0:
            issues.append("含 %d 个表格" % len(tables))
        if max(levels) if levels else 0:
            if max(levels) >= 3:
                issues.append("存在 %d 级深缩进，层级偏深" % (max(levels) + 1))

        role = guess_role(idx, len(slides), texts, title, tables, image_count)
        if role != "content":
            issues.append("页面角色：%s" % {
                "cover": "封面（写开场钩子）",
                "toc": "目录（预告全篇结构）",
                "section": "章节扉页（写转场衔接）",
                "ending": "结束页（回顾 + 行动号召）",
                "image": "纯图片页（结合画面说明，不得臆造）",
            }[role])

        pages.append({
            "index": idx,
            "role": role,
            "title": title,
            "texts": texts,
            "tables": tables,
            "imageCount": image_count,
            "notes": notes,
            "charCount": page_chars,
            "infoDensity": density,
            "issues": issues,
        })

    # 文档级问题
    doc_issues = []
    if len(font_set) > 2:
        doc_issues.append("字体不统一，全篇出现 %d 种：%s"
                          % (len(font_set), "、".join(sorted(font_set))))
    if font_sizes:
        min_size = min(font_sizes)
        if min_size < 12:
            doc_issues.append("存在小于 12pt 的小字，最小 %.1fpt" % min_size)
    else:
        min_size = None
        doc_issues.append("未检测到显式字号（可能全部继承主题）")
    if has_notes == 0:
        doc_issues.append("全篇无讲解备注")
    elif has_notes < len(slides):
        doc_issues.append("仅 %d/%d 页有备注" % (has_notes, len(slides)))

    result = {
        "ok": True,
        "source": {
            "file": os.path.abspath(path),
            "sha1": sha1_of(path),
            "slideCount": len(slides),
            "canvas": {"widthPt": emu_to_pt(prs.slide_width),
                       "heightPt": emu_to_pt(prs.slide_height)},
        },
        "docStats": {
            "totalChars": total_chars,
            "fontSet": sorted(font_set),
            "minFontPt": min_size,
            "notesCount": has_notes,
            "noTextPages": no_text_pages,
            "issues": doc_issues,
        },
        "pages": pages,
    }
    return result, (no_text_pages == len(slides))


def main():
    ap = argparse.ArgumentParser(description="解析 PPTX → page-content.json")
    ap.add_argument("input", help="原始课件 .pptx")
    ap.add_argument("-o", "--output", default=None, help="输出 JSON 路径")
    ap.add_argument("-w", "--workdir", default="work", help="工作目录（默认 work/）")
    args = ap.parse_args()

    out = args.output or os.path.join(args.workdir, "page-content.json")
    try:
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    except Exception as e:
        fail(EXIT_BAD_INPUT, "输出目录无法创建：%s" % e,
             "请更换输出路径，确认盘符与目录存在且有写权限")

    result, all_no_text = analyze(args.input)

    with open(out, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    s = result["source"]
    d = result["docStats"]
    sys.stderr.write("[analyze] 解析完成：%d 页 / 总字数 %d / 有备注 %d 页\n"
                     % (s["slideCount"], d["totalChars"], d["notesCount"]))
    for it in d["issues"]:
        sys.stderr.write("  · %s\n" % it)
    sys.stderr.write("[analyze] 已输出：%s\n" % out)

    # 全文无文本：不阻塞，但用 6 提示模型按「纯图片页」规则处理
    sys.exit(EXIT_NO_TEXT if all_no_text else EXIT_OK)


if __name__ == "__main__":
    main()
