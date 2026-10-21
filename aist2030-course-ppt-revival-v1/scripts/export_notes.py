#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
export_notes.py —— 老旧课程 PPT 复活 · 第五步「导出备注摘要」

从（优化后的）PPTX 读取每页备注栏，导出按页整理的讲解备注摘要，
可直接交付讲师，也可直接喂给下游 AIST2030-mass-ppt-narration-video-v3 出片。

用法：
    python export_notes.py 优化后.pptx -o 讲解备注摘要.md          # Markdown（默认）
    python export_notes.py 优化后.pptx -o 讲稿.txt --format txt     # 纯文本

退出码：
    0 成功 | 2 文件问题 | 3 损坏 | 4 加密/旧格式 | 5 空文档
    6 全篇无备注（不阻塞，仍产出摘要并列出空备注页） | 8 写出失败
"""

import argparse
import json
import os
import sys

try:
    from pptx import Presentation
except ImportError:
    sys.stderr.write("[export] 缺少 python-pptx，请先 pip install python-pptx\n")
    sys.exit(1)

EXIT_OK = 0
EXIT_BAD_INPUT = 2
EXIT_PARSE_FAILED = 3
EXIT_ENCRYPTED = 4
EXIT_EMPTY = 5
EXIT_NO_NOTES = 6
EXIT_WRITE_FAILED = 8

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# 与下游出片环节实测语速一致：约 4.8 字/秒
CHARS_PER_SEC = 4.8


def fail(code, message, hint=""):
    sys.stderr.write(json.dumps(
        {"ok": False, "code": code, "message": message, "hint": hint},
        ensure_ascii=False, indent=2) + "\n")
    sys.exit(code)


def fmt_mmss(seconds):
    m = int(seconds // 60)
    s = int(round(seconds % 60))
    if s == 60:
        m, s = m + 1, 0
    return "%d:%02d" % (m, s)


def collect(path):
    if not os.path.exists(path):
        fail(EXIT_BAD_INPUT, "文件不存在：%s" % path, "请确认路径")
    if os.path.splitext(path)[1].lower() != ".pptx":
        fail(EXIT_BAD_INPUT, "仅支持 .pptx：%s" % path, "请先另存为 .pptx")
    with open(path, "rb") as f:
        if f.read(8) == OLE_MAGIC:
            fail(EXIT_ENCRYPTED, "文件是加密包或旧式二进制格式", "请去除密码后另存为 .pptx")

    try:
        prs = Presentation(path)
    except Exception as e:
        fail(EXIT_PARSE_FAILED, "文件损坏或无法解析：%s" % e, "请另存为新的 .pptx 后重试")

    slides = list(prs.slides)
    if not slides:
        fail(EXIT_EMPTY, "该文档不含任何幻灯片", "请确认文件")

    pages = []
    for idx, slide in enumerate(slides, start=1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = slide.shapes.title.text.strip()
        except Exception:
            pass

        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
        except Exception:
            notes = ""

        pages.append({
            "index": idx,
            "title": title or "（无标题）",
            "notes": notes,
            "chars": len(notes.replace("\n", "")),
        })
    return pages


def render_md(pages, src_name):
    total = sum(p["chars"] for p in pages)
    empty = [p["index"] for p in pages if not p["notes"]]
    lines = [
        "# 讲解备注摘要 · %s" % src_name,
        "",
        "- 页数：%d" % len(pages),
        "- 备注总字数：%d 字" % total,
        "- 预估讲解时长：约 %s（按 4.8 字/秒折算）" % fmt_mmss(total / CHARS_PER_SEC),
    ]
    if empty:
        lines.append("- ⚠️ 以下页面备注栏为空：%s" % "、".join(str(i) for i in empty))
    lines.append("")
    lines.append("---")
    lines.append("")

    for p in pages:
        lines.append("## 第 %d 页 · %s" % (p["index"], p["title"]))
        lines.append("")
        if p["notes"]:
            lines.append("> %s" % p["notes"].replace("\n", "\n> "))
            lines.append("")
            lines.append("（本页 %d 字，约 %s）" % (p["chars"], fmt_mmss(p["chars"] / CHARS_PER_SEC)))
        else:
            lines.append("（本页无备注，需补充）")
        lines.append("")
    return "\n".join(lines)


def render_txt(pages, src_name):
    total = sum(p["chars"] for p in pages)
    lines = [
        "讲解备注摘要 · %s" % src_name,
        "页数：%d / 备注总字数：%d / 预估时长：约 %s"
        % (len(pages), total, fmt_mmss(total / CHARS_PER_SEC)),
        "",
    ]
    for p in pages:
        lines.append("【第 %d 页】%s" % (p["index"], p["title"]))
        lines.append(p["notes"] if p["notes"] else "（本页无备注，需补充）")
        lines.append("")
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser(description="导出按页讲解备注摘要")
    ap.add_argument("input", help="（优化后的）PPTX")
    ap.add_argument("-o", "--output", default=None, help="输出路径")
    ap.add_argument("--format", default="md", choices=["md", "txt"], help="摘要格式")
    ap.add_argument("-w", "--workdir", default="work", help="工作目录")
    args = ap.parse_args()

    pages = collect(args.input)
    src_name = os.path.basename(args.input)

    ext = "md" if args.format == "md" else "txt"
    out = args.output or os.path.join(args.workdir, "notes-summary.%s" % ext)
    try:
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    except Exception as e:
        fail(EXIT_WRITE_FAILED, "输出目录无法创建：%s" % e,
             "请更换输出路径，确认盘符与目录存在且有写权限")

    body = render_md(pages, src_name) if args.format == "md" else render_txt(pages, src_name)
    try:
        with open(out, "w", encoding="utf-8") as f:
            f.write(body)
    except Exception as e:
        fail(EXIT_WRITE_FAILED, "摘要写出失败：%s" % e, "请更换输出路径")

    total = sum(p["chars"] for p in pages)
    empty = [p["index"] for p in pages if not p["notes"]]
    sys.stderr.write("[export] 已导出：%s\n" % out)
    sys.stderr.write("[export] %d 页 / %d 字 / 预估讲解 %s\n"
                     % (len(pages), total, fmt_mmss(total / CHARS_PER_SEC)))
    if empty:
        sys.stderr.write("[export] ⚠️ 以下页备注为空：%s\n" % "、".join(str(i) for i in empty))

    sys.exit(EXIT_NO_NOTES if not total else EXIT_OK)


if __name__ == "__main__":
    main()
